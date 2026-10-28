# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook(r"2026年年度计划汇总.xlsx", data_only=True)
ws = wb['2026年量值溯源总表']
for r in range(1, ws.max_row+1):
    val = str(ws.cell(r, 4).value or '') # Col 4 is equipment ID
    if val.startswith('WJ'):
        print(f"Row {r}: Code={val}, Name={ws.cell(r, 3).value}, Group={ws.cell(r, 2).value}")
