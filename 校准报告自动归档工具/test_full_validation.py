# -*- coding: utf-8 -*-
import os
import sys
import shutil
import datetime
import openpyxl

from main_gui import (
    build_archive_dir,
    build_new_filename,
    parse_date_obj,
    add_years_minus_one_day,
    sync_ledger_with_records,
    LedgerDatabase,
    VERSION
)

def test_all():
    print(f"=== 🚀 开始全量功能自动化验证 (版本: {VERSION}) ===")

    # 1. 验证归档目录生成器 (按年月与组别)
    print("\n--- 1. 验证归档目录生成逻辑 ---")
    fields_quant = {
        'asset_no': 'H0406',
        'inst_name': 'pH计',
        'cal_date': '20260728',
        'group': '理化微生物组',
        'device_type': '定量'
    }
    
    # 1.1 源目录就地分类 (organize 模式)
    src_dir = r"D:\下载报告\2026年7-9月达丰校准证书"
    dir_org = build_archive_dir(None, 'DAF', fields_quant, source_dir=src_dir, is_organize_in_source=True)
    expected_leaf = "2026年7-9月达丰校准证书-理化微生物组"
    assert os.path.basename(dir_org) == expected_leaf, f"预期 {expected_leaf}, 实际 {os.path.basename(dir_org)}"
    print(f"✅ organize 模式目录测试通过: {dir_org}")

    # 1.2 归档至目标总库模式
    archive_root = r"D:\工作\2026年检定校准证书"
    dir_archive = build_archive_dir(archive_root, 'DAF', fields_quant, source_dir=src_dir, is_organize_in_source=False)
    expected_leaf2 = "2026年7月达丰校准证书-理化微生物组"
    assert os.path.basename(dir_archive) == expected_leaf2, f"预期 {expected_leaf2}, 实际 {os.path.basename(dir_archive)}"
    print(f"✅ copy/move 目标库目录测试通过: {dir_archive}")

    # 2. 验证台账载入与信息提取
    print("\n--- 2. 验证台账数据库载入 ---")
    ledger = LedgerDatabase()
    ledger.load()
    quant_path = ledger.find_ledger_path(is_quant=True)
    quick_path = ledger.find_ledger_path(is_quant=False)
    print(f"✅ 定量台账路径: {quant_path}")
    print(f"✅ 快检台账路径: {quick_path}")
    assert 'H0406' in ledger.quantitative_map, "H0406 应在 quantitative_map 中"
    info_h0406 = ledger.quantitative_map['H0406']
    print(f"✅ H0406 台账信息: group={info_h0406.get('group')}, last_cal={info_h0406.get('last_cal_date')}, cycle={info_h0406.get('cycle')}")

    # 3. 验证定量台账双向回填更新
    print("\n--- 3. 验证定量台账双向回填更新 ---")
    test_quant_copy = "test_quant_plan_copy.xlsx"
    shutil.copyfile(quant_path, test_quant_copy)
    
    records_quant = [
        {'设备编号': 'H0406', '校准日期': '2026-07-28', '机构': '达丰', '仪器名称': 'pH计'},
        {'设备编号': 'H0187', '校准日期': '2026-05-15', '机构': '达丰', '仪器名称': '三重四级杆-气质联用仪'}
    ]
    res_q = sync_ledger_with_records(test_quant_copy, records_quant, is_quant=True)
    assert res_q['status'] == 'success', f"更新失败: {res_q}"
    print(f"✅ 定量更新成功: 更新处数={res_q['updated_count']}, 备份文件={os.path.basename(res_q['backup_path'])}")
    assert os.path.exists(res_q['backup_path']), "备份文件必须真实存在"
    
    # 验证公式是否保留
    wb_q = openpyxl.load_workbook(test_quant_copy, data_only=False)
    ws_trace = wb_q['2026年量值溯源总表']
    # 查找 H0406 行
    h0406_row = None
    for r in range(2, ws_trace.max_row+1):
        if ws_trace.cell(r, 3).value == 'H0406':
            h0406_row = r
            break
    assert h0406_row is not None
    i_val = ws_trace.cell(h0406_row, 9).value
    k_formula = ws_trace.cell(h0406_row, 11).value
    l_note = ws_trace.cell(h0406_row, 12).value
    print(f"✅ H0406 更新后核对: Col9(上次校准)={i_val}, Col11(公式)={k_formula}, Col12(备注)={l_note}")
    assert i_val == datetime.datetime(2026, 7, 28, 0, 0)
    assert str(k_formula).startswith('='), "Col 11 公式必须完整保留"
    assert '2026已校准' in str(l_note)
    wb_q.close()
    
    # 清理定量测试文件与备份
    os.remove(test_quant_copy)
    if os.path.exists(res_q['backup_path']):
        os.remove(res_q['backup_path'])

    # 4. 验证快检台账双向回填更新
    print("\n--- 4. 验证快检台账双向回填更新 ---")
    test_quick_copy = "test_quick_plan_copy.xlsx"
    shutil.copyfile(quick_path, test_quick_copy)
    
    records_quick = [
        {'设备编号': '22110001', '校准日期': '2026-01-20', '机构': '帝恩', '仪器名称': '农药残留快速检测仪'},
        {'设备编号': 'BB031', '校准日期': '2026-07-25', '机构': '帝恩', '仪器名称': '高通量农药残留分析仪'}
    ]
    res_k = sync_ledger_with_records(test_quick_copy, records_quick, is_quant=False)
    assert res_k['status'] == 'success', f"快检更新失败: {res_k}"
    print(f"✅ 快检更新成功: 更新处数={res_k['updated_count']}, 备份文件={os.path.basename(res_k['backup_path'])}")
    assert os.path.exists(res_k['backup_path']), "快检备份文件必须真实存在"
    
    os.remove(test_quick_copy)
    if os.path.exists(res_k['backup_path']):
        os.remove(res_k['backup_path'])

    print("\n🎉🎉 全量测试全部通过！双向更新、组别归档、公式防护、安全备份功能100%健全！")

if __name__ == '__main__':
    test_all()
